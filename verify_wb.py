import warnings, pickle, sys; sys.path.insert(0,'.'); warnings.filterwarnings('ignore')
import openpyxl, pandas as pd, numpy as np
from src.config import CAPITAL, OUT, ROOT
wb=openpyxl.load_workbook(ROOT/'FINESSE_CITADEL_SUBMISSION.xlsx')
print('OPENS OK. Sheets:'); [print('  -',s) for s in wb.sheetnames]
names=set(wb.sheetnames)
D=pickle.load(open(OUT/'frozen_run.pkl','rb')); rep=D['IS']['report']; nav=D['IS']['nav']
bad=0; nf=0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith('='):
                nf+=1
                for tok in c.value.split("'"):
                    pass
                import re
                for m in re.findall(r"'([^']+)'!", c.value):
                    if m not in names: print('BAD SHEETREF',ws.title,c.coordinate,m); bad+=1
                for m in re.findall(r"(?<![A-Za-z'])([A-Za-z][A-Za-z0-9_. &]*)!", c.value):
                    if m not in names and m.strip() not in names: print('BAD SHEETREF',ws.title,c.coordinate,m); bad+=1
print(f'\n{nf} live formulas found, {bad} broken sheet references')
# independent value cross-checks (what Excel WILL compute)
n=len(nav); last=5+n
r=nav.pct_change().dropna()
yrs=(nav.index[-1]-nav.index[0]).days/365.25
cagr=(nav.iloc[-1]/CAPITAL)**(1/yrs)-1
sd=r.std(ddof=1)
peak=nav.cummax(); dd=(nav/peak-1).min()
print('\nCROSS-CHECK (Excel formula target vs Python report):')
print(f"  Total Net PNL   Excel=B{last}-B3 -> {nav.iloc[-1]-CAPITAL:,.0f}   python={rep['total_net_pnl']:,.0f}")
print(f"  CAGR            -> {cagr*100:.4f}%   python={rep['annualised_return_pct']:.4f}%")
print(f"  MDD             -> {dd*100:.4f}%   python={rep['max_drawdown_pct']:.4f}%")
print(f"  SD daily        -> {sd:.6f}")
print(f"  Sharpe (comp)   -> {cagr/sd:.4f}   python={rep['sharpe_competition']:.4f}")
print(f"  Sharpe (std)    -> {(r.mean()*252)/(sd*np.sqrt(252)):.4f}   python={rep['sharpe_standard']:.4f}")
tr=D['IS']['trades']
print(f"  Trade Log cost  -> live sum {(tr['notional']*0.001).sum():,.2f}   booked {tr['cost'].sum():,.2f}")
# constraint proofs
tw=D['IS']['tw']
mx=max(len(v) for v in tw.values()); sums=[abs(v.sum()-1) for v in tw.values()]; mw=max(v.max() for v in tw.values())
print(f'\nCONSTRAINT PROOF: max concurrent names={mx} (limit 10) | max |sum(w)-1|={max(sums):.2e} | max single weight={mw:.4f} (cap 0.20)')
